"""Tests for the multi-evaluator merge mode of report.py.

Covers: two-input merge produces the reference wide layout (two-row header,
per-evaluator columns), merged statistics computed on per-query means,
attribution restricted to merged-WA badcases with all-N WA text,
review-sheet stacking, alignment validation, and single-input compatibility.
"""

import os
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

import sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

from report import main as report_main

EN_HEADERS = [
    "query", "标注", "filename", "搜索结果", "score", "大模型问答结果",
    "Response Relevance-Score", "Response Relevance-Reason",
    "Citation Accuracy-Score", "Citation Accuracy-Reason",
    "Scope Completeness-Score", "Scope Completeness-Reason",
    "Instruction Following-Score", "Instruction Following-Reason",
    "Readability and Naturalness-Score", "Readability and Naturalness-Reason",
    "Weighted Average", "Overall Summary",
]

QUERY_ROWS = [
    {
        "query": "What is a p-frame?",
        "annotation": '[{"chunk_id": "Chunk 8", "source_text": "A pframe can be defined as a semi-fixed multi-word expression."}]',
        "filename": "Pho_2010.pdf",
        "retrieval": "Type 3 p-frames are those that occurred in two or more rhetorical move-steps.",
        "score": 0.54,
        "answer": "A p-frame can be categorized into three types: Type 1, Type 2, Type 3.",
    },
    {
        "query": "A good query?",
        "annotation": '[{"chunk_id": "Chunk 1", "source_text": "The answer is clearly yes."}]',
        "filename": "GoodDoc.pdf",
        "retrieval": "The answer is clearly yes and supported by evidence.",
        "score": 0.9,
        "answer": "Yes, clearly.",
    },
]

REVIEW_HEADERS = ["行号", "Query", "RR 原分", "最终结论"]


def _dims(wa: float):
    """Deterministic dim set whose weighted average equals wa."""
    import math
    for a in range(1, 6):
        for b in range(1, 6):
            for c in range(1, 6):
                for d in range(1, 6):
                    for e in range(1, 6):
                        if abs(0.3 * a + 0.3 * b + 0.15 * c + 0.15 * d + 0.1 * e - wa) < 1e-9:
                            return [a, b, c, d, e]
    raise ValueError(f"no dim set for wa={wa}")


def _build_input(path: str, wa1: float, wa2: float, review_rows: list[list] | None = None,
                 query_rows=None, review_headers=None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation"
    ws.append(EN_HEADERS)
    rows = query_rows if query_rows is not None else QUERY_ROWS
    for base, wa in zip(rows, [wa1, wa2]):
        dims = _dims(wa)
        reasons = ["reason%d" % (i + 1) for i in range(5)]
        ws.append([
            base["query"], base["annotation"], base["filename"], base["retrieval"],
            base["score"], base["answer"],
            dims[0], reasons[0], dims[1], reasons[1], dims[2], reasons[2],
            dims[3], reasons[3], dims[4], reasons[4],
            wa, "summary of %s" % base["query"][:8],
        ])
    if review_rows is not None:
        rws = wb.create_sheet("Review")
        headers = review_headers if review_headers is not None else REVIEW_HEADERS
        rws.append(headers)
        for row in review_rows:
            rws.append(row)
    wb.save(path)


class MergeModeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="kb_merge_")
        cls.a = os.path.join(cls.tmp, "evaluator_a.xlsx")
        cls.b = os.path.join(cls.tmp, "evaluator_b.xlsx")
        # wa 2.2 and 3.4 -> merged 2.8 (badcase); 4.0 and 4.6 -> merged 4.3
        _build_input(cls.a, 2.2, 4.0, review_rows=[[0, "qA", "3", "manual_review"]])
        _build_input(cls.b, 3.4, 4.6, review_rows=[[1, "qB", "4", "needs_adjustment"],
                                                   [0, "qA", "5", "reasonable"]])
        cls.out = os.path.join(cls.tmp, "merged_report.xlsx")
        report_main([
            "--input", cls.a, "--input", cls.b,
            "--output", cls.out,
            "--language", "en",
            "--no-llm",
        ])
        cls.wb = load_workbook(cls.out)

    def test_sheet_names_and_order(self):
        self.assertEqual(self.wb.sheetnames,
                         ["Evaluation", "低分复核", "统计摘要", "Badcase归因分析"])

    def test_evaluation_two_row_header(self):
        ws = self.wb["Evaluation"]
        group = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        column = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(group[0], "公共信息")
        self.assertEqual(group[6], "Response Relevance（回答相关性）")
        self.assertEqual(group[26], "加权均分")
        self.assertEqual(group[28], "综合评价")
        self.assertEqual(column[6], "分数_1")
        self.assertEqual(column[7], "理由_1")
        self.assertEqual(column[8], "分数_2")
        self.assertEqual(column[9], "理由_2")
        self.assertEqual(column[26], "加权均分_1")
        self.assertEqual(column[27], "加权均分_2")
        self.assertEqual(column[28], "综合评价_1")
        self.assertEqual(column[29], "综合评价_2")
        self.assertEqual(ws.max_column, 6 + 4 * 5 + 4)  # 6 public + 20 dim + 2 wa + 2 summary

    def test_evaluation_group_headers_have_distinct_colors(self):
        ws = self.wb["Evaluation"]
        # Top-left cell of each merged group: 公共信息(A1) + 5 dims + 加权均分 + 综合评价
        group_anchors = [1, 7, 11, 15, 19, 23, 27, 29]
        colors = {}
        for c in group_anchors:
            fill = ws.cell(row=1, column=c).fill
            self.assertEqual(fill.fill_type, "solid", f"group anchor c{c} has no fill")
            colors.setdefault(fill.fgColor.rgb, []).append(ws.cell(row=1, column=c).value)
        self.assertEqual(len(colors), 8)
        # 加权均分 dark blue, 综合评价 near-black
        self.assertEqual(ws.cell(row=1, column=27).fill.fgColor.rgb, "002F5597")
        self.assertEqual(ws.cell(row=1, column=29).fill.fgColor.rgb, "00262626")
        # column header row uses a lighter tint than its group color
        tint = ws.cell(row=2, column=7).fill
        group = ws.cell(row=1, column=7).fill
        self.assertEqual(tint.fill_type, "solid")
        self.assertNotEqual(tint.fgColor.rgb, group.fgColor.rgb)

    def test_column_header_dark_text_on_light_tint(self):
        ws = self.wb["Evaluation"]
        font = ws.cell(row=2, column=7).font
        self.assertEqual(font.color.rgb, "001F1F1F")  # dark text, not white
        # group row keeps white bold on the dark palette
        group_font = ws.cell(row=1, column=7).font
        self.assertEqual(group_font.color.rgb, "00FFFFFF")

    def test_evaluation_data_public_cols_preserved(self):
        ws = self.wb["Evaluation"]
        self.assertEqual(ws.cell(row=3, column=1).value, "What is a p-frame?")
        self.assertEqual(ws.cell(row=4, column=1).value, "A good query?")
        self.assertEqual(ws.max_row, 4)

    def test_evaluation_data_evaluator_cells_under_own_headers(self):
        """Regression: data fill must follow the dimension-major header layout.

        Each dimension group (cols 7-10, 11-14, ...) holds evaluator 1's
        分数/理由 pair then evaluator 2's; WA group is cols 27-28, summary
        group cols 29-30. A previous evaluator-major fill shifted every
        evaluator block sideways under the wrong headers.
        """
        ws = self.wb["Evaluation"]
        dims_a1, dims_a2 = _dims(2.2), _dims(4.0)
        dims_b1, dims_b2 = _dims(3.4), _dims(4.6)
        # Row 3 = query 1, dim1 group: evaluator 1 pair then evaluator 2 pair.
        self.assertEqual(ws.cell(row=3, column=7).value, float(dims_a1[0]))
        self.assertEqual(ws.cell(row=3, column=8).value, "reason1")
        self.assertEqual(ws.cell(row=3, column=9).value, float(dims_b1[0]))
        self.assertEqual(ws.cell(row=3, column=10).value, "reason1")
        # dim2 group.
        self.assertEqual(ws.cell(row=3, column=11).value, float(dims_a1[1]))
        self.assertEqual(ws.cell(row=3, column=12).value, "reason2")
        self.assertEqual(ws.cell(row=3, column=13).value, float(dims_b1[1]))
        self.assertEqual(ws.cell(row=3, column=14).value, "reason2")
        # WA group: each evaluator's recomputed WA for that row.
        self.assertEqual(ws.cell(row=3, column=27).value, 2.2)
        self.assertEqual(ws.cell(row=3, column=28).value, 3.4)
        self.assertIn("summary of", str(ws.cell(row=3, column=29).value))
        self.assertIn("summary of", str(ws.cell(row=3, column=30).value))
        # Row 4 = query 2.
        self.assertEqual(ws.cell(row=4, column=7).value, float(dims_a2[0]))
        self.assertEqual(ws.cell(row=4, column=9).value, float(dims_b2[0]))
        self.assertEqual(ws.cell(row=4, column=27).value, 4.0)
        self.assertEqual(ws.cell(row=4, column=28).value, 4.6)

    def test_summary_merged_means(self):
        summary = self.wb["统计摘要"]
        self.assertIn("2 评测员合并", str(summary["A1"].value))
        self.assertEqual(summary["A6"].value, 3.55)  # (2.8 + 4.3) / 2
        self.assertEqual(summary["C6"].value, "1 / 2")  # only merged wa 2.8 < 3.0
        # 序号 = raw row index in the first evaluator's data (0-based),
        # same semantics as the single-input report.
        self.assertEqual(summary["A27"].value, 0)
        self.assertEqual(summary["B27"].value, 2.8)
        self.assertEqual(summary["C27"].value, "2.20 / 3.40")
        self.assertEqual(summary["D27"].value, "What is a p-frame?")

    def test_attribution_only_merged_badcases_with_all_was(self):
        attribution = self.wb["Badcase归因分析"]
        self.assertEqual(attribution["A3"].value, 1)
        self.assertEqual(attribution["B3"].value, "What is a p-frame?")
        self.assertEqual(attribution["H3"].value, "2.20 / 3.40")
        self.assertIn("归因口径", str(attribution.cell(row=4, column=1).value))
        self.assertEqual(attribution.cell(row=5, column=1).value, None)  # only one badcase row

    def test_review_sheet_stacked_column_union(self):
        review = self.wb["低分复核"]
        headers = [review.cell(row=1, column=c).value for c in range(1, review.max_column + 1)]
        self.assertEqual(headers, REVIEW_HEADERS)
        self.assertEqual(review.max_row, 4)  # header + a(1) + b(2)
        self.assertEqual(review.cell(row=2, column=4).value, "manual_review")
        self.assertEqual(review.cell(row=4, column=4).value, "reasonable")


class MergeValidationTest(unittest.TestCase):
    def test_mismatched_queries_refused(self):
        tmp = tempfile.mkdtemp(prefix="kb_merge_bad_")
        a = os.path.join(tmp, "a.xlsx")
        b = os.path.join(tmp, "b.xlsx")
        _build_input(a, 2.2, 4.0)
        different = [dict(QUERY_ROWS[0]), dict(QUERY_ROWS[1])]
        different[1]["query"] = "A different query?"
        _build_input(b, 3.4, 4.6, query_rows=different)
        out = os.path.join(tmp, "out.xlsx")
        with self.assertRaises(SystemExit) as ctx:
            report_main(["--input", a, "--input", b, "--output", out,
                         "--language", "en", "--no-llm"])
        self.assertIn("Cannot merge", str(ctx.exception))

    def test_mismatched_raw_layout_refused(self):
        """Same queries/answers but different raw row layouts must be refused.

        The merged Evaluation sheet fills evaluator cells by raw-row position,
        so a block-layout file (continuation rows) merged with a flat file
        would silently shift one evaluator's scores onto the wrong queries.
        """
        tmp = tempfile.mkdtemp(prefix="kb_merge_layout_")
        a = os.path.join(tmp, "a.xlsx")
        b = os.path.join(tmp, "b.xlsx")
        _build_input(a, 2.2, 4.0)  # flat: 2 raw data rows
        # Evaluator b scores the same two queries but uses a block layout
        # (one continuation row under query 1) -> 3 raw data rows.
        wb = Workbook()
        ws = wb.active
        ws.title = "Evaluation"
        ws.append(EN_HEADERS)
        dims = _dims(3.4)
        reasons = ["reason%d" % (i + 1) for i in range(5)]
        for qi, wa in ((0, 3.4), (1, 4.6)):
            base = QUERY_ROWS[qi]
            ws.append([
                base["query"], base["annotation"], base["filename"],
                base["retrieval"], base["score"], base["answer"],
                dims[0], reasons[0], dims[1], reasons[1], dims[2], reasons[2],
                dims[3], reasons[3], dims[4], reasons[4],
                wa, "summary of %s" % base["query"][:8],
            ])
            if qi == 0:
                ws.append(["", "", "extra.pdf", "extra chunk", 0.5] + [""] * 13)
        wb.save(b)
        out = os.path.join(tmp, "out.xlsx")
        with self.assertRaises(SystemExit) as ctx:
            report_main(["--input", a, "--input", b, "--output", out,
                         "--language", "en", "--no-llm"])
        self.assertIn("Cannot merge", str(ctx.exception))
        self.assertIn("raw row count", str(ctx.exception))


class SingleInputCompatTest(unittest.TestCase):
    def test_single_input_still_single_format(self):
        tmp = tempfile.mkdtemp(prefix="kb_merge_single_")
        a = os.path.join(tmp, "a.xlsx")
        out = os.path.join(tmp, "out.xlsx")
        _build_input(a, 2.2, 4.0)
        report_main(["--input", a, "--output", out, "--language", "en", "--no-llm"])
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Evaluation", "统计摘要", "Badcase归因分析"])
        self.assertEqual(wb["统计摘要"]["A6"].value, 3.1)  # (2.2 + 4.0) / 2 single WA


if __name__ == "__main__":
    unittest.main()
