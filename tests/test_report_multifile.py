"""Tests for multi-file attribution in report.py.

Covers the fix: badcase 归因/原文核验 must traverse every recalled file and
every chunk of a block, not just the first row. Block/grouped layout inputs
(query only on block header rows) are aggregated per query group: 语料源文件
shows all dedup filenames, 实际召回段落 aggregates all chunk texts, and
read_from_index receives the full filename list for 原文核验.
"""

import os
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

import sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

from report import main as report_main

EN_HEADERS = [
    "query", "标注", "filename", "搜索结果", "score", "大模型问答结果",
    "Response Relevance-Score", "Response Relevance-Reason",
    "Citation Accuracy-Score", "Citation Accuracy-Reason",
    "Scope Completeness-Score", "Scope Completeness-Reason",
    "Instruction Following-Score", "Instruction Following-Reason",
    "Readability and Naturalness-Score", "Readability and Naturalness-Reason",
    "Weighted Average", "Overall Summary",
]

# Dim-score sets whose recomputed weighted averages are exactly 2.2
# (badcase) and 4.8 (good) under WEIGHTS (0.3, 0.3, 0.15, 0.15, 0.1).
_DIMS_BAD = [2, 2, 2, 2, 4]
_DIMS_GOOD = [5, 5, 5, 5, 3]


def _block_input(path: str, prefix: str = "") -> None:
    """Block/grouped-layout evaluated workbook.

    Query 1 spans 3 chunk rows across 2 files (fileA.pdf twice, fileB.pdf
    once) -> filenames dedup to [fileA, fileB]; wa 2.2 -> badcase.
    Query 2 is a single-row block; wa 4.8 (not a badcase).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation"
    ws.append(EN_HEADERS)
    # Block header row: query + annotation + first chunk + answer + scores.
    # Annotation "beta golden" only appears in fileB's chunk (row 2), so a
    # first-row-only implementation would report 未召回; aggregation must hit 部分命中.
    # Score/reason columns are interleaved (Score, Reason, Score, Reason, ...)
    # exactly like EN_HEADERS — writing all scores first would shift every
    # reason column and zero out half the dims.
    ws.append([
        f"{prefix}What is a p-frame?",
        "beta golden",
        "fileA.pdf",
        "chunk text alpha one",
        0.9,
        "Answer says p-frame is X.",
        _DIMS_BAD[0], "reason1", _DIMS_BAD[1], "reason2", _DIMS_BAD[2], "reason3",
        _DIMS_BAD[3], "reason4", _DIMS_BAD[4], "reason5",
        2.2, "bad summary",
    ])
    # Continuation row 1: another file's chunk (no query/answer/scores).
    ws.append(["", "", "fileB.pdf", "chunk text beta two", 0.8] + [""] * 13)
    # Continuation row 2: same file as header (dedup) + third chunk.
    ws.append(["", "", "fileA.pdf", "chunk text alpha three", 0.7] + [""] * 13)
    # Second block: single row, good score.
    ws.append([
        "A good query?",
        "gold good",
        "good.pdf",
        "good chunk text",
        0.95,
        "Yes clearly.",
        _DIMS_GOOD[0], "g1", _DIMS_GOOD[1], "g2", _DIMS_GOOD[2], "g3",
        _DIMS_GOOD[3], "g4", _DIMS_GOOD[4], "g5",
        4.8, "good summary",
    ])
    wb.save(path)


class _FakeApi:
    base_url = "http://fake"
    api_key = "fake-key"


class _FakeSettings:
    base_url = "http://fake"
    api_key = "fake-key"


class MultiFileAttributionTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="kb_multifile_")
        # Corpus dir with stem-match docs (fileA.pdf -> fileA.md, fileB.pdf -> fileB.txt).
        cls.doc_dir = os.path.join(cls.tmp, "docs")
        os.makedirs(cls.doc_dir)
        with open(os.path.join(cls.doc_dir, "fileA.md"), "w", encoding="utf-8") as f:
            f.write("CONTENT-ALPHA-UNIQUE\n" * 5)
        with open(os.path.join(cls.doc_dir, "fileB.txt"), "w", encoding="utf-8") as f:
            f.write("CONTENT-BETA-UNIQUE\n" * 5)

        cls.input_path = os.path.join(cls.tmp, "block_evaluated.xlsx")
        _block_input(cls.input_path)

        cls.output_path = os.path.join(cls.tmp, "block_report.xlsx")
        report_main([
            "--input", cls.input_path,
            "--output", cls.output_path,
            "--language", "en",
            "--no-llm",
        ])
        cls.wb = load_workbook(cls.output_path)

    def test_attribution_source_column_lists_all_files(self):
        attr = self.wb["Badcase归因分析"]
        # Badcase = query 1; column C (语料源文件) shows dedup, ordered filenames.
        self.assertEqual(attr["C3"].value, "fileA.pdf, fileB.pdf")
        # WA must be the recomputed 2.20 of _DIMS_BAD — guards the fixture's
        # interleaved score/reason column order.
        self.assertEqual(attr["H3"].value, "2.20")

    def test_attribution_retrieval_column_aggregates_all_chunks(self):
        attr = self.wb["Badcase归因分析"]
        retrieval = attr["E3"].value
        for chunk in ("chunk text alpha one", "chunk text beta two", "chunk text alpha three"):
            self.assertIn(chunk, retrieval, f"retrieval missing chunk: {chunk}")

    def test_attribution_coverage_uses_aggregated_retrieval(self):
        # gold "beta" appears only in fileB's chunk; coverage must be computed
        # from the full block text => 部分命中, not a first-row-only 未召回.
        attr = self.wb["Badcase归因分析"]
        coverage = attr["F3"].value
        self.assertIn("覆盖率", coverage)
        self.assertIn("部分命中", coverage)

    def test_single_row_block_still_single_file(self):
        # Query 2 is a good case (not a badcase) -> not in attribution sheet;
        # but ensure it was counted and did not corrupt the block aggregation.
        summary = self.wb["统计摘要"]
        self.assertIn("2", str(summary["C6"].value))  # 1 badcase / 2 total


class MultiFileLLMTest(unittest.TestCase):
    """原文核验 path: read_from_index must receive every recalled filename."""

    def test_original_text_contains_all_files(self):
        tmp = tempfile.mkdtemp(prefix="kb_multifile_llm_")
        doc_dir = os.path.join(tmp, "docs")
        os.makedirs(doc_dir)
        with open(os.path.join(doc_dir, "fileA.md"), "w", encoding="utf-8") as f:
            f.write("CONTENT-ALPHA-UNIQUE " * 10)
        with open(os.path.join(doc_dir, "fileB.txt"), "w", encoding="utf-8") as f:
            f.write("CONTENT-BETA-UNIQUE " * 10)

        input_path = os.path.join(tmp, "in.xlsx")
        _block_input(input_path, prefix="LLM ")
        out_path = os.path.join(tmp, "out.xlsx")

        with mock.patch("report.load_api_settings", return_value=_FakeSettings()), \
             mock.patch("report.call_attribution_model", return_value={
                 "归因结论": "A 召回问题（纯）",
                 "归因分析": "分析",
                 "原文核验": "✅ 已核实",
             }) as api_mock:
            report_main([
                "--input", input_path,
                "--output", out_path,
                "--language", "en",
                "--model", "fake-model",
                "--doc-dirs", doc_dir,
            ])

        # The attribution prompt for the badcase must embed both files' text.
        prompts = [call.args[0] for call in api_mock.call_args_list]
        self.assertEqual(len(prompts), 1)
        prompt = prompts[0]
        self.assertIn("CONTENT-ALPHA-UNIQUE", prompt)
        self.assertIn("CONTENT-BETA-UNIQUE", prompt)
        self.assertIn("fileA.md", prompt)
        self.assertIn("fileB.txt", prompt)


class MultiFileMergedTest(unittest.TestCase):
    """Merged mode: filenames list must flow through _merged_rows."""

    def test_merged_attribution_lists_all_files(self):
        tmp = tempfile.mkdtemp(prefix="kb_multifile_merge_")
        a = os.path.join(tmp, "a.xlsx")
        b = os.path.join(tmp, "b.xlsx")
        _block_input(a)
        _block_input(b)
        out = os.path.join(tmp, "merged.xlsx")
        report_main(["--input", a, "--input", b, "--output", out,
                     "--language", "en", "--no-llm"])
        wb = load_workbook(out)
        attr = wb["Badcase归因分析"]
        self.assertEqual(attr["C3"].value, "fileA.pdf, fileB.pdf")
        retrieval = attr["E3"].value
        self.assertIn("chunk text beta two", retrieval)
        # Both evaluators recomputed the same badcase WA.
        self.assertEqual(attr["H3"].value, "2.20 / 2.20")


if __name__ == "__main__":
    unittest.main()