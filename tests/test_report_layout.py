"""Structural golden test: generated report layout must match the reference
workbook (tests/fixtures/reference_layout.xlsx) for sheet names, merged
ranges, and header rows, while data rows follow the same section rules.
"""

import os
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

import sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

from report import main as report_main

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")
REFERENCE = os.path.join(FIXTURES_DIR, "reference_layout.xlsx")

EN_HEADERS = [
    "query", "标注", "filename", "搜索结果", "score", "大模型问答结果",
    "Response Relevance-Score", "Response Relevance-Reason",
    "Citation Accuracy-Score", "Citation Accuracy-Reason",
    "Scope Completeness-Score", "Scope Completeness-Reason",
    "Instruction Following-Score", "Instruction Following-Reason",
    "Readability and Naturalness-Score", "Readability and Naturalness-Reason",
    "Weighted Average", "Overall Summary",
]


def _build_input_workbook(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation"
    ws.append(EN_HEADERS)
    ws.append([
        "What is a p-frame?",
        '[{"chunk_id": "Chunk 8", "source_text": "A pframe can be defined as a semi-fixed multi-word expression."}]',
        "Pho_2010.pdf",
        "Type 3 p-frames are those that occurred in two or more rhetorical move-steps.",
        0.54,
        "A p-frame can be categorized into three types: Type 1, Type 2, Type 3.",
        2, "Answered classification instead of definition",
        2, "Definition not supported by retrieved text",
        2, "Missing the core definition",
        2, "User asked for a definition, not a taxonomy",
        4, "Clear sentences",
        2.0, "Low score summary",
    ])
    ws.append([
        "A good query?",
        '[{"chunk_id": "Chunk 1", "source_text": "The answer is clearly yes."}]',
        "GoodDoc.pdf",
        "The answer is clearly yes and supported by evidence.",
        0.9,
        "Yes, clearly.",
        5, "Hits the core intent",
        5, "Fully supported",
        4, "Covers the main points",
        5, "Direct format",
        4, "Readable",
        4.8, "Good summary",
    ])
    wb.save(path)


class ReportLayoutTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp_dir = tempfile.mkdtemp(prefix="kb_report_")
        cls.input_path = os.path.join(cls.tmp_dir, "input_evaluated.xlsx")
        cls.output_path = os.path.join(cls.tmp_dir, "input_report.xlsx")
        _build_input_workbook(cls.input_path)
        report_main([
            "--input", cls.input_path,
            "--output", cls.output_path,
            "--language", "en",
            "--no-llm",
        ])
        cls.wb = load_workbook(cls.output_path)

    def test_sheet_names_and_order(self):
        self.assertEqual(self.wb.sheetnames, ["Evaluation", "统计摘要", "Badcase归因分析"])

    def test_summary_merged_ranges_match_reference(self):
        reference = load_workbook(REFERENCE)["统计摘要"]
        summary = self.wb["统计摘要"]
        static_reference = {
            str(rng) for rng in reference.merged_cells.ranges
            if str(rng) not in {"C27:F27", "C28:F28", "C29:F29", "C30:F30", "C31:F31", "C32:F32"}
        }
        actual = {str(rng) for rng in summary.merged_cells.ranges}
        self.assertTrue(static_reference.issubset(actual), f"missing static merges: {static_reference - actual}")
        self.assertIn("C26:F26", actual)
        self.assertIn("C27:F27", actual)

    def test_summary_header_rows_match_reference(self):
        reference = load_workbook(REFERENCE)["统计摘要"]
        summary = self.wb["统计摘要"]
        for row in (5, 10, 18, 26):
            expected = [reference.cell(row=row, column=c).value for c in range(1, 7)]
            actual = [summary.cell(row=row, column=c).value for c in range(1, 7)]
            self.assertEqual(actual, expected, f"header row {row} mismatch")

    def test_summary_section_titles(self):
        summary = self.wb["统计摘要"]
        self.assertTrue(str(summary["A1"].value).startswith("EN 知识库问答评测 · 统计摘要"))
        self.assertEqual(summary["A4"].value, "核心指标")
        self.assertEqual(summary["A9"].value, "维度评分明细")
        self.assertEqual(summary["A17"].value, "分数段分布（按加权均分）")
        self.assertIn("Badcase 明细", summary["A25"].value)

    def test_summary_core_indicators(self):
        summary = self.wb["统计摘要"]
        self.assertEqual(summary["A6"].value, 3.48)  # (2.2 + 4.75) / 2
        self.assertEqual(summary["C6"].value, "1 / 2")
        self.assertIn("达标", str(summary["D6"].value))

    def test_summary_band_rows(self):
        summary = self.wb["统计摘要"]
        bands = [summary.cell(row=r, column=1).value for r in range(19, 24)]
        self.assertEqual(bands, ["1.0-1.9", "2.0-2.9", "3.0-3.9", "4.0-4.5", "4.5-5.0"])

    def test_summary_badcase_rows(self):
        summary = self.wb["统计摘要"]
        self.assertEqual(summary["A27"].value, 0)
        self.assertEqual(summary["B27"].value, 2.2)

    def test_attribution_title_merged_and_headers(self):
        attribution = self.wb["Badcase归因分析"]
        reference = load_workbook(REFERENCE)["Badcase归因分析"]
        self.assertIn("A1:L1", [str(rng) for rng in attribution.merged_cells.ranges])
        expected_headers = [reference.cell(row=2, column=c).value for c in range(1, 13)]
        actual_headers = [attribution.cell(row=2, column=c).value for c in range(1, 13)]
        self.assertEqual(actual_headers, expected_headers)
        self.assertIn("Badcase 归因分析", attribution["A1"].value)

    def test_attribution_rows_and_footnote(self):
        attribution = self.wb["Badcase归因分析"]
        self.assertEqual(attribution["A3"].value, 1)
        self.assertEqual(attribution["B3"].value, "What is a p-frame?")
        self.assertIn("覆盖率", attribution["F3"].value)
        self.assertIn("A 召回问题（纯）", attribution["I3"].value)
        footnote = attribution.cell(row=4, column=1).value
        self.assertIn("归因口径", footnote)
        self.assertIn("A 纯召回 1 条", footnote)

    def test_input_sheet_preserved(self):
        eval_ws = self.wb["Evaluation"]
        self.assertEqual(eval_ws.max_row, 3)
        self.assertEqual(eval_ws.cell(row=2, column=1).value, "What is a p-frame?")


if __name__ == "__main__":
    unittest.main()
