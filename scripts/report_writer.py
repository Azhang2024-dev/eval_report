"""Report sheets (统计摘要 + Badcase归因分析) in the reference layout.

The layout mirrors ``en_llm_sampled_100_evaluated_human_2_v2.xlsx``:
the summary sheet uses the fixed section/merged-cell structure of that file,
and the attribution sheet uses its 12-column table with a merged title row
and a merged attribution-criteria footnote row.
"""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill

import pandas as pd

from scripts.excel_style import HEADER_FONT, HEADER_FILL, HEADER_ALIGN, auto_width

_SECTION_FONT = Font(bold=True, size=12, color="1F4E79")
_TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
_GROUP_FONT = Font(bold=True, size=11, color="FFFFFF")
# Column-header row sits on light group tints -> dark text for readability.
_COL_FONT = Font(bold=True, size=10, color="1F1F1F")

# Canonical Chinese dimension names per the kb-qa-eval prompt standard.
_DIM_CN = {
    "Response Relevance": "回答相关性",
    "Citation Accuracy": "引用准确性",
    "Scope Completeness": "范围完整性",
    "Instruction Following": "指令遵循",
    "Readability and Naturalness": "阅读感受与自然度",
}

# Distinct group colors, all chosen dark enough for white bold text
# (contrast ratio >= ~4.5:1).
_GROUP_COLORS = {
    "公共信息": "595959",
    "回答相关性": "4472C4",
    "引用准确性": "548235",
    "范围完整性": "C55A11",
    "指令遵循": "7030A0",
    "阅读感受与自然度": "00838F",
    "加权均分": "2F5597",
    "综合评价": "262626",
}

# Light tint of the group color for the column-header row (dark text).
_GROUP_TINTS = {
    "公共信息": "D9D9D9",
    "回答相关性": "BDD7EE",
    "引用准确性": "E2EFDA",
    "范围完整性": "FCE4D6",
    "指令遵循": "CCC0DA",
    "阅读感受与自然度": "E0F2F1",
    "加权均分": "B4C6E7",
    "综合评价": "E7E6E6",
}


def bilingual_dim_label(label: str) -> str:
    """Append the canonical Chinese name to an English dimension label.

    ``Response Relevance`` -> ``Response Relevance（回答相关性）``.
    Labels already containing CJK text are returned unchanged.
    """
    if label in _DIM_CN:
        return f"{label}（{_DIM_CN[label]}）"
    return label


def _group_color_key(label: str) -> str:
    """Map a (possibly bilingual) group label to its color key.

    ``Response Relevance（回答相关性）`` -> ``回答相关性``; bare labels such as
    公共信息 / 加权均分 pass through unchanged.
    """
    if "（" in label and label.endswith("）"):
        return label.split("（", 1)[1][:-1]
    return label

DIM_HEADERS = ["维度", "均值", "中位数", "标准差", "最低", "最高"]
BAND_HEADERS = ["分数段", "数量", "占比"]
BADCASE_HEADERS = ["序号", "加权均分", "查询"]
MERGED_BADCASE_HEADERS = ["序号", "合并均分", "WA（评测员1/…/N）", "查询"]

_PUBLIC_COL_ORDER = ["query", "reference", "filename", "chunks", "score", "answer"]


def _fmt_threshold(value: float) -> str:
    return f"{value:.1f}" if value == int(value) else f"{value:g}"

ATTRIBUTION_HEADERS = [
    "序号",
    "Query",
    "语料源文件",
    "金标答案段落（标注）",
    "实际召回段落（搜索结果）",
    "金标↔召回匹配",
    "大模型回答",
    "WA（本文件 / 对侧文件）",
    "归因结论",
    "归因分析",
    "原文核验",
    "评分Reason摘录",
]

MOS_PASS = 4.1
BADCASE_RATE_PASS = 10.0


def _style_headers(ws, row: int, headers: list[str], start_col: int = 1) -> None:
    for c, header in enumerate(headers, start_col):
        cell = ws.cell(row=row, column=c, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def write_summary_sheet(
    wb,
    *,
    language_code: str,
    total: int,
    empty: int,
    failed: int,
    dim_labels: list[str],
    dim_stats: list[dict[str, float]],
    weighted: list[float],
    mos: float,
    badcase_rate: float,
    badcases: list[tuple[int, float, str]],
    threshold: float = 3.0,
) -> Any:
    """Write the 统计摘要 sheet; returns the worksheet."""
    ws = wb.create_sheet("统计摘要")
    lang = language_code.upper()
    threshold_text = _fmt_threshold(threshold)

    ws.merge_cells("A1:F1")
    ws["A1"] = f"{lang} 知识库问答评测 · 统计摘要"
    ws["A1"].font = _TITLE_FONT

    ws.merge_cells("A2:F2")
    ws["A2"] = f"评测总数：{total} 条 query    回答为空 / 失败：{empty} / {failed}"

    ws.merge_cells("A4:E4")
    ws["A4"] = "核心指标"
    ws["A4"].font = _SECTION_FONT

    summary_headers = [
        "MOS（加权均分）",
        f"Badcase 率（<{threshold_text}）",
        "Badcase 数",
        "MOS 准出（≥4.1）",
        "Badcase 准出（≤10%）",
    ]
    _style_headers(ws, 5, summary_headers)
    mos_pass = "✅ 达标" if mos >= MOS_PASS else "❌ 未达标"
    rate_pass = "✅ 达标" if badcase_rate <= BADCASE_RATE_PASS else "❌ 未达标"
    ws["A6"] = round(mos, 2)
    ws["B6"] = f"{badcase_rate:.1f}%"
    ws["C6"] = f"{len(badcases)} / {total}"
    ws["D6"] = mos_pass
    ws["E6"] = rate_pass

    ws.merge_cells("A7:E7")
    ws["A7"] = "准出标准：MOS ≥ 4.1 且 Badcase 率 ≤ 10%"

    ws.merge_cells("A9:F9")
    ws["A9"] = "维度评分明细"
    ws["A9"].font = _SECTION_FONT

    _style_headers(ws, 10, DIM_HEADERS)
    for i, (label, stat) in enumerate(zip(dim_labels, dim_stats)):
        r = 11 + i
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=round(stat["mean"], 2))
        ws.cell(row=r, column=3, value=round(stat["median"], 1))
        ws.cell(row=r, column=4, value=round(stat["stdev"], 2))
        ws.cell(row=r, column=5, value=stat["min"])
        ws.cell(row=r, column=6, value=stat["max"])

    r = 17
    ws.merge_cells(f"A{r}:F{r}")
    ws.cell(row=r, column=1, value="分数段分布（按加权均分）").font = _SECTION_FONT
    _style_headers(ws, r + 1, BAND_HEADERS)
    from scripts.eval_stats import band_distribution
    bands = band_distribution(weighted)
    for i, (band, count) in enumerate(bands.items()):
        row = r + 2 + i
        pct = count / len(weighted) * 100 if weighted else 0
        ws.cell(row=row, column=1, value=band)
        ws.cell(row=row, column=2, value=count)
        ws.cell(row=row, column=3, value=f"{pct:.1f}%")

    r = 25
    ws.merge_cells(f"A{r}:F{r}")
    ws.cell(row=r, column=1, value=f"Badcase 明细（加权均分 < {threshold_text}，共 {len(badcases)} 条）").font = _SECTION_FONT
    ws.merge_cells(f"C{r + 1}:F{r + 1}")
    _style_headers(ws, r + 1, BADCASE_HEADERS)
    for i, (index, wa, query) in enumerate(badcases):
        row = r + 2 + i
        ws.cell(row=row, column=1, value=index)
        ws.cell(row=row, column=2, value=round(wa, 2))
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        ws.cell(row=row, column=3, value=query)

    auto_width(ws, min_width=12)
    return ws


def write_attribution_sheet(
    wb,
    *,
    language_code: str,
    model_label: str,
    threshold: float,
    rows: list[dict[str, Any]],
    criteria_footnote: str,
) -> Any:
    """Write the Badcase归因分析 sheet; returns the worksheet."""
    ws = wb.create_sheet("Badcase归因分析")
    lang = language_code.upper()
    if model_label and "评测员" in model_label:
        model_part = f"（{model_label}）"
    else:
        model_part = f"（评测模型 {model_label}）" if model_label else ""
    ws.merge_cells("A1:L1")
    ws["A1"] = (
        f"Badcase 归因分析 · {lang}{model_part} · {len(rows)} 条"
        f"（Weighted Average < {threshold}）· 判定方法：金标标注 vs 实际召回客观匹配 + 原文回溯核验"
    )

    _style_headers(ws, 2, ATTRIBUTION_HEADERS)

    for i, row in enumerate(rows):
        r = 3 + i
        values = [
            i + 1,
            row.get("query", ""),
            row.get("filename", ""),
            row.get("annotation", ""),
            row.get("retrieval", ""),
            row.get("coverage_text", ""),
            row.get("answer", ""),
            row.get("wa_text", ""),
            row.get("verdict", ""),
            row.get("analysis", ""),
            row.get("verification", ""),
            row.get("reasons", ""),
        ]
        for c, value in enumerate(values, 1):
            ws.cell(row=r, column=c, value=value)

    last_row = 3 + len(rows)
    ws.merge_cells(start_row=last_row, start_column=1, end_row=last_row, end_column=12)
    ws.cell(row=last_row, column=1, value=criteria_footnote)
    ws.cell(row=last_row, column=1).alignment = HEADER_ALIGN

    widths = {"A": 6, "B": 36, "C": 28, "D": 44, "E": 44, "F": 30, "G": 44, "H": 14, "I": 20, "J": 40, "K": 40, "L": 44}
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    for row_cells in ws.iter_rows(min_row=3, max_row=3 + len(rows)):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def _num_or_blank(value) -> Any:
    """Return the value if it is a positive number, else blank string."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    return number if number > 0 else ""


def write_merged_evaluation_sheet(wb, loaded: list[dict[str, Any]]) -> Any:
    """Write the merged multi-evaluator Evaluation sheet (two-row header).

    Columns: 6 public + per dimension group: 分数_k/理由_k for every evaluator
    k in 1..N, followed by 加权均分_1..N and 综合评价_1..N. Data rows mirror
    the first evaluator's raw layout; evaluator cells are filled per raw-row
    index (identical raw layouts are validated upstream).

    The column layout is described ONCE as 'col_defs' and drives both the
    header rendering and the data fill, so the two can never diverge. WA
    cells show the recomputed weighted averages — the same numbers the
    统计摘要 reports — rather than the input files' (possibly stale) WA column.
    """
    first = loaded[0]
    n = len(loaded)
    raw_df = first["raw_df"]
    source_cols = first["source_cols"]
    score_pairs = first["score_pairs"]
    dim_labels = [bilingual_dim_label(pair[0].rsplit("-", 1)[0]) for pair in score_pairs]

    ws = wb.create_sheet("Evaluation")

    public_headers = [source_cols.get(field, "") for field in _PUBLIC_COL_ORDER]

    # Dimension-major layout: each dimension group holds every evaluator's
    # 分数_k/理由_k pair, then the WA group and the summary group follow.
    # Each entry is ("public", field) | ("score", k, d) | ("reason", k, d)
    # | ("wa", k) | ("summary", k).
    col_defs: list[tuple] = [("public", field) for field in _PUBLIC_COL_ORDER]
    group_span: list[tuple[str, int, int]] = [("公共信息", 1, len(_PUBLIC_COL_ORDER))]
    col = len(_PUBLIC_COL_ORDER) + 1
    for d, label in enumerate(dim_labels):
        start = col
        for k in range(n):
            col_defs.append(("score", k, d))
            col_defs.append(("reason", k, d))
            col += 2
        group_span.append((label, start, col - 1))
    wa_start = col
    for k in range(n):
        col_defs.append(("wa", k))
        col += 1
    group_span.append(("加权均分", wa_start, col - 1))
    summary_start = col
    for k in range(n):
        col_defs.append(("summary", k))
        col += 1
    group_span.append(("综合评价", summary_start, col - 1))

    for c, spec in enumerate(col_defs, 1):
        kind = spec[0]
        if kind == "public":
            group = "公共信息"
            column_header = public_headers[c - 1]
        elif kind in ("score", "reason"):
            group = dim_labels[spec[2]]
            column_header = ("分数" if kind == "score" else "理由") + f"_{spec[1] + 1}"
        elif kind == "wa":
            group = "加权均分"
            column_header = f"加权均分_{spec[1] + 1}"
        else:
            group = "综合评价"
            column_header = f"综合评价_{spec[1] + 1}"
        color_key = _group_color_key(group)
        cell = ws.cell(row=1, column=c, value=group)
        cell.font = _GROUP_FONT
        cell.fill = PatternFill("solid", fgColor=_GROUP_COLORS.get(color_key, "4472C4"))
        cell.alignment = HEADER_ALIGN
        col_cell = ws.cell(row=2, column=c, value=column_header)
        col_cell.font = _COL_FONT
        col_cell.fill = PatternFill("solid", fgColor=_GROUP_TINTS.get(color_key, "D9E2F3"))
        col_cell.alignment = HEADER_ALIGN
    for group, start, end in group_span:
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)

    score_cols = [score for score, _reason in score_pairs]
    reason_cols = [reason for _score, reason in score_pairs]
    per_evaluator = []
    for entry in loaded:
        per_evaluator.append({
            "df": entry["raw_df"],
            "summary_header": entry["summary_header"],
            # Recomputed WA keyed by raw block-header row index.
            "wa_by_raw": {row["index"]: row["wa"] for row in entry["rows"]},
        })

    for i, (_, row) in enumerate(raw_df.iterrows()):
        r = i + 3
        e_rows = [ev["df"].iloc[i] if i < len(ev["df"]) else None for ev in per_evaluator]
        for c, spec in enumerate(col_defs, 1):
            kind = spec[0]
            if kind == "public":
                column = source_cols.get(spec[1])
                value = row.get(column) if column else None
            else:
                evaluator = per_evaluator[spec[1]]
                e_row = e_rows[spec[1]]
                if kind == "score":
                    value = _num_or_blank(e_row.get(score_cols[spec[2]])) if e_row is not None else ""
                elif kind == "reason":
                    value = e_row.get(reason_cols[spec[2]]) if e_row is not None else None
                elif kind == "wa":
                    wa = evaluator["wa_by_raw"].get(i)
                    value = round(wa, 2) if wa is not None else None
                else:  # summary
                    value = e_row.get(evaluator["summary_header"]) if e_row is not None else None
            ws.cell(row=r, column=c, value=value)

    auto_width(ws, min_width=10)
    return ws


def write_merged_review_sheet(wb, loaded: list[dict[str, Any]]) -> Any:
    """Stack every evaluator's review sheet (column union, missing filled)."""
    frames = []
    for entry in loaded:
        review_df = entry.get("review_df")
        if review_df is None or review_df.empty:
            continue
        header = [str(h) if h is not None else "" for h in review_df.iloc[0].tolist()]
        data = review_df.iloc[1:].reset_index(drop=True)
        data.columns = header
        frames.append((header, data))
    if not frames:
        return None
    union: list[str] = []
    for header, _data in frames:
        for column in header:
            if column and column not in union:
                union.append(column)

    ws = wb.create_sheet("低分复核")
    for c, column in enumerate(union, 1):
        cell = ws.cell(row=1, column=c, value=column)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
    r = 2
    for _header, data in frames:
        for _, row in data.iterrows():
            for c, column in enumerate(union, 1):
                value = row.get(column)
                ws.cell(row=r, column=c,
                        value=None if value is None or (isinstance(value, float) and pd.isna(value)) else value)
            r += 1
    auto_width(ws, min_width=10)
    return ws


def write_merged_summary_sheet(
    wb,
    *,
    language_code: str,
    evaluator_count: int,
    total: int,
    empty: int,
    failed: int,
    dim_labels: list[str],
    dim_stats: list[dict[str, float]],
    weighted: list[float],
    mos: float,
    badcase_rate: float,
    badcases: list[tuple[int, float, str]],
    merged_rows: list[dict[str, Any]],
    threshold: float = 3.0,
) -> Any:
    """Write the 统计摘要 sheet for a merged multi-evaluator report."""
    ws = wb.create_sheet("统计摘要")
    lang = language_code.upper()
    threshold_text = _fmt_threshold(threshold)
    wa_by_index = {row["index"]: row for row in merged_rows}

    ws.merge_cells("A1:F1")
    ws["A1"] = f"{lang} 知识库问答评测 · 统计摘要（{evaluator_count} 评测员合并）"
    ws["A1"].font = _TITLE_FONT

    ws.merge_cells("A2:F2")
    ws["A2"] = (f"评测总数：{total} 条 query（{evaluator_count} 评测员合并均分）    "
                f"Badcase 判定：合并均分 < {threshold_text}")

    ws.merge_cells("A4:E4")
    ws["A4"] = "核心指标"
    ws["A4"].font = _SECTION_FONT

    summary_headers = [
        "MOS（合并均分）",
        f"Badcase 率（<{threshold_text}）",
        "Badcase 数",
        "MOS 准出（≥4.1）",
        "Badcase 准出（≤10%）",
    ]
    _style_headers(ws, 5, summary_headers)
    mos_pass = "✅ 达标" if mos >= MOS_PASS else "❌ 未达标"
    rate_pass = "✅ 达标" if badcase_rate <= BADCASE_RATE_PASS else "❌ 未达标"
    ws["A6"] = round(mos, 2)
    ws["B6"] = f"{badcase_rate:.1f}%"
    ws["C6"] = f"{len(badcases)} / {total}"
    ws["D6"] = mos_pass
    ws["E6"] = rate_pass

    ws.merge_cells("A7:E7")
    ws["A7"] = "准出标准：MOS ≥ 4.1 且 Badcase 率 ≤ 10%"

    ws.merge_cells("A9:F9")
    ws["A9"] = "维度评分明细（评测员分数取均值后统计）"
    ws["A9"].font = _SECTION_FONT

    _style_headers(ws, 10, DIM_HEADERS)
    for i, (label, stat) in enumerate(zip(dim_labels, dim_stats)):
        r = 11 + i
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=round(stat["mean"], 2))
        ws.cell(row=r, column=3, value=round(stat["median"], 1))
        ws.cell(row=r, column=4, value=round(stat["stdev"], 2))
        ws.cell(row=r, column=5, value=round(stat["min"], 2))
        ws.cell(row=r, column=6, value=round(stat["max"], 2))

    r = 17
    ws.merge_cells(f"A{r}:F{r}")
    ws.cell(row=r, column=1, value="分数段分布（按合并均分）").font = _SECTION_FONT
    _style_headers(ws, r + 1, BAND_HEADERS)
    from scripts.eval_stats import band_distribution
    bands = band_distribution(weighted)
    for i, (band, count) in enumerate(bands.items()):
        row = r + 2 + i
        pct = count / len(weighted) * 100 if weighted else 0
        ws.cell(row=row, column=1, value=band)
        ws.cell(row=row, column=2, value=count)
        ws.cell(row=row, column=3, value=f"{pct:.1f}%")

    r = 25
    ws.merge_cells(f"A{r}:F{r}")
    ws.cell(row=r, column=1, value=f"Badcase 明细（合并均分 < {threshold_text}，共 {len(badcases)} 条）").font = _SECTION_FONT
    ws.merge_cells(f"D{r + 1}:F{r + 1}")
    _style_headers(ws, r + 1, MERGED_BADCASE_HEADERS)
    for i, (index, wa, query) in enumerate(badcases):
        row = r + 2 + i
        # 序号 = raw row index in the first evaluator's data (0-based),
        # consistent with the single-input report.
        ws.cell(row=row, column=1, value=index)
        ws.cell(row=row, column=2, value=round(wa, 2))
        merged = wa_by_index.get(index, {})
        was = merged.get("was") or []
        ws.cell(row=row, column=3,
                value=" / ".join(f"{w:.2f}" if w is not None else "-" for w in was))
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
        ws.cell(row=row, column=4, value=query)

    auto_width(ws, min_width=12)
    return ws
