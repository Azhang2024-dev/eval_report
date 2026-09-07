"""Original-document reading for the 原文核验 column of the attribution report.

Mirrors kb-qa-eval's review_module behavior: text files are read directly,
pdfs via pymupdf text layer, docx via python-docx, xlsx via openpyxl,
anything else becomes a ``[Binary file: ...]`` placeholder.

When a gold annotation is supplied, injection is relevance-ranked: every
source file is scored by gold-segment bigram coverage (reusing coverage.py),
files are injected highest-score first, and an oversized high-score file is
excerpted by best-matching paragraphs instead of blind head truncation — so
the gold sentence stays in the material even when it sits in a late file or
a late paragraph.
"""

from __future__ import annotations

import functools
import os
import re


def _stem(name: str) -> str:
    """Lowercased basename without its extension, e.g. 'a.PDF' -> 'a'."""
    return os.path.splitext(os.path.basename(name))[0].lower()


def find_doc(doc_dirs: list[str], filename: str) -> str | None:
    """Locate *filename* (basename or relative path) inside *doc_dirs*.

    Exact basename match wins; otherwise the stem (name without extension)
    is matched so a report referencing ``x.pdf`` resolves to ``x.md`` in the
    cleaned corpus (same matching rule as kb-qa-eval's review_module).
    Resolution goes through build_index so all reading paths apply the same
    rules, including the ambiguity guard on the stem fallback.
    """
    if not filename:
        return None
    index = build_index(doc_dirs)
    basename = os.path.basename(str(filename).strip())
    path = index.get(basename.lower())
    if path is None:
        path = index.get(_stem(filename))
    return path


def build_index(doc_dirs: list[str]) -> dict[str, str]:
    """Walk *doc_dirs* once and return ``{basename.lower(): path}`` plus stem
    fallback entries ``{stem: path}`` (only when the stem is unambiguous)."""
    index: dict[str, str] = {}
    stem_count: dict[str, set[str]] = {}
    for directory in doc_dirs or []:
        for root, _dirs, files in os.walk(directory):
            for name in files:
                index.setdefault(name.lower(), os.path.join(root, name))
                stem = _stem(name)
                stem_count.setdefault(stem, set()).add(name.lower())
    for stem, names in stem_count.items():
        if len(names) == 1:
            index.setdefault(stem, index[list(names)[0]])
    return index


def _resolve_path(index: dict[str, str], filename: str) -> str | None:
    basename = os.path.basename(str(filename).strip())
    return index.get(basename.lower()) or index.get(_stem(filename))


# Bigram-overlap noise floor: below this, a file's gold "coverage" is random
# character overlap, not signal, so it scores 0 and keeps its original order.
RANK_MIN_SCORE = 0.1

# Paragraph excerpt caps: a truncated file contributes its first paragraph
# plus at most EXCERPT_MAX_PARAGRAPHS top-scoring paragraphs (each scoring at
# least EXCERPT_MIN_SCORE), so one saturated giant file cannot swallow the
# whole budget with weakly-overlapping filler.
EXCERPT_MAX_PARAGRAPHS = 12
EXCERPT_MIN_SCORE = 0.15


def rank_filenames_by_gold(
    index: dict[str, str], filenames: list[str], annotation
) -> list[tuple[str, float]]:
    """Score each filename by gold coverage and sort descending (stable).

    The file score is the max coverage of any gold segment inside the file's
    full text, using the same character-bigram metric as the 金标↔召回匹配
    column. Files that do not resolve, or gold-less annotations, score 0 and
    keep their original relative order (Python sort is stable).
    """
    from coverage import parse_gold_segments, segment_coverage

    segments = parse_gold_segments(annotation)
    if not segments or not filenames:
        return [(f, 0.0) for f in (filenames or [])]
    scored: list[tuple[str, float]] = []
    for filename in filenames:
        path = _resolve_path(index, filename)
        if path is None:
            scored.append((filename, 0.0))
            continue
        text = read_doc(path)
        if not text or text.startswith("[Binary file:"):
            scored.append((filename, 0.0))
            continue
        score = max(segment_coverage(seg, text) for seg in segments)
        if score < RANK_MIN_SCORE:
            score = 0.0
        scored.append((filename, round(score, 4)))
    scored.sort(key=lambda pair: pair[1], reverse=True)
    return scored


def _excerpt_by_gold(text: str, max_content: int, annotation) -> str:
    """Excerpt an oversized document by gold-relevant paragraphs.

    Paragraphs (blank-line separated) are scored by gold-segment coverage;
    the top-scoring paragraphs (at most ``EXCERPT_MAX_PARAGRAPHS``, each at
    or above ``EXCERPT_MIN_SCORE``) are kept in original document order, and
    the first paragraph is always retained (titles/metadata). The result is
    bounded by *max_content*. Falls back to head truncation when there is no
    gold or no paragraph matches it.
    """
    from coverage import parse_gold_segments, segment_coverage

    segments = parse_gold_segments(annotation)
    if not segments or max_content <= 0:
        return text[:max_content]
    paragraphs = [p for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(paragraphs) <= 1:
        return text[:max_content]
    para_scores = [
        (max(segment_coverage(seg, para) for seg in segments), i)
        for i, para in enumerate(paragraphs)
    ]
    if all(score <= 0 for score, _ in para_scores):
        return text[:max_content]
    chosen = {0}
    used = len(paragraphs[0])
    picked = 0
    for score, i in sorted(para_scores, key=lambda t: (-t[0], t[1])):
        if score < EXCERPT_MIN_SCORE or picked >= EXCERPT_MAX_PARAGRAPHS:
            break
        if i in chosen:
            picked += 1
            continue
        cost = len(paragraphs[i]) + 2  # paragraph + joiner
        if used + cost <= max_content:
            chosen.add(i)
            used += cost
            picked += 1
    result = "\n\n".join(paragraphs[i] for i in sorted(chosen))
    if len(result) > max_content:
        result = result[:max_content]
    return result


def _read_with_budget(
    index: dict[str, str],
    filenames: list[str],
    budget: int = 128000,
    annotation=None,
) -> str:
    """Shared budget loop for read_from_index / read_docs_with_budget.

    With *annotation*, files are injected in gold-coverage order, headers
    carry the objective coverage label, and truncated files are excerpted by
    gold-relevant paragraphs instead of blind head truncation. Both the
    per-file header and the content count against *budget*; once the
    remaining budget cannot fit the next header, stop.
    """
    if annotation is not None:
        ranked = rank_filenames_by_gold(index, filenames, annotation)
        scores = dict(ranked)
        ordered = [f for f, _ in ranked]
    else:
        scores = None
        ordered = list(filenames or [])
    parts: list[str] = []
    used = 0
    for filename in ordered:
        basename = os.path.basename(str(filename).strip())
        path = _resolve_path(index, filename)
        if path is None:
            parts.append(f"- {basename}: (not found in doc dirs)")
            continue
        score = scores.get(filename, 0.0) if scores is not None else None
        if score is not None and score > 0:
            header = f"- {os.path.basename(path)}（金标覆盖 {score * 100:.0f}%）:\n"
        else:
            header = f"- {os.path.basename(path)}:\n"
        remaining = budget - used
        if remaining <= len(header):
            break
        text = read_doc(path)
        max_content = remaining - len(header)
        if len(text) > max_content:
            if annotation is not None:
                text = _excerpt_by_gold(text, max_content, annotation)
            else:
                text = text[:max_content]
        parts.append(header + text)
        used += len(header) + len(text)
    return "\n\n".join(parts)


def read_from_index(
    index: dict[str, str], filenames: list[str], budget: int = 128000, annotation=None
) -> str:
    """Read the original docs for *filenames* via a prebuilt index.

    With ``annotation`` provided, injection is relevance-ranked (see
    ``rank_filenames_by_gold`` / ``_excerpt_by_gold``); without it the
    historical file-order injection is preserved exactly.
    """
    return _read_with_budget(index, filenames, budget, annotation)


@functools.lru_cache(maxsize=128)
def read_doc(path: str) -> str:
    """Read a document into plain text with binary placeholders.

    Cached: corpus files are immutable within a run and the same source file
    is often referenced by multiple badcase rows (lru_cache avoids repeating
    expensive pdf/docx/xlsx extraction).
    """
    ext = os.path.splitext(path)[1].lower()
    name = os.path.basename(path)
    try:
        if ext in (".md", ".txt", ".text"):
            with open(path, "r", encoding="utf-8", errors="replace") as fh:
                return fh.read()
        if ext == ".pdf":
            import fitz
            try:
                doc = fitz.open(path)
                text = "\n".join(page.get_text() for page in doc)
                doc.close()
                return text or f"[Binary file: {name} (no text layer)]"
            except Exception:
                return f"[Binary file: {name}]"
        if ext == ".docx":
            import docx
            document = docx.Document(path)
            return "\n".join(p.text for p in document.paragraphs) or f"[Binary file: {name} (no extractable text)]"
        if ext == ".xlsx":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=True)
                blocks = []
                for sheet in wb.worksheets:
                    lines = [
                        " | ".join("" if v is None else str(v).strip() for v in row)
                        for row in sheet.iter_rows(values_only=True)
                        if any(v is not None for v in row)
                    ]
                    if lines:
                        blocks.append(f"=== Sheet: {sheet.title} ===\n" + "\n".join(lines))
                wb.close()
                return "\n\n".join(blocks) or f"[Binary file: {name} (no extractable text)]"
            except Exception:
                return f"[Binary file: {name}]"
    except Exception:
        pass
    return f"[Binary file: {name}]"


def read_docs_with_budget(
    doc_dirs: list[str], filenames: list[str], budget: int = 128000, annotation=None
) -> str:
    """Read the original docs for *filenames*, truncating to *budget* chars.

    Builds the corpus index once and applies the same resolution, ranking,
    and budget accounting as read_from_index.
    """
    index = build_index(doc_dirs)
    return _read_with_budget(index, filenames, budget, annotation)
